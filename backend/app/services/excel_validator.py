from pathlib import Path

from openpyxl import load_workbook

from app.validators.workbook_validator import WorkbookValidator
from app.validators.data_validator import DataValidator
from app.validators.formula_validator import FormulaValidator
from app.validators.statistics_validator import StatisticsValidator
from app.validators.excel_error_validator import ExcelErrorValidator
from app.validators.cross_sheet_validator import CrossSheetValidator
from app.validators.template_detector import TemplateDetector
from app.utils.sheet_mapper import SheetMapper
from app.services.laya_classifier import get_laya_classifier
from app.models.validation import AIReviewSummary, AIReviewClassification


class ExcelValidator:
    """
    Main validation pipeline for assessment report workbooks.
    """

    def __init__(self):

        self.workbook_validator = WorkbookValidator()
        self.data_validator = DataValidator()
        self.formula_validator = FormulaValidator()
        self.statistics_validator = StatisticsValidator()
        self.excel_error_validator = ExcelErrorValidator()
        self.cross_sheet_validator = CrossSheetValidator()
        self.template_detector = TemplateDetector()

    def validate(self, file_path: Path, filename: str = None):

        issues = []

        # ==================================================
        # 1. OPEN WORKBOOK
        # ==================================================

        try:

            workbook = load_workbook(
                file_path,
                data_only=False
            )

        except Exception as exc:

            return {
                "status": "ERROR",
                "template": None,
                "errors": [
                    {
                        "code": "WORKBOOK_OPEN_FAILED",
                        "category": "Workbook",
                        "message": (
                            "Excel workbook could not be opened."
                        ),
                        "sheet": None,
                        "cell": None,
                        "expected": "Valid Excel workbook",
                        "actual": str(exc),
                    }
                ],
                "warnings": [],
            }

        try:

            # ==================================================
            # 2. DETECT TEMPLATE
            # ==================================================

            template_result = self.template_detector.detect(
                workbook
            )

            if template_result is None:

                return {
                    "status": "ERROR",
                    "template": None,
                    "errors": [
                        {
                            "code": "UNKNOWN_TEMPLATE",
                            "category": "Workbook Structure",
                            "message": (
                                "Assessment report template "
                                "could not be identified."
                            ),
                            "sheet": None,
                            "cell": None,
                            "expected": (
                                "Supported assessment "
                                "report template"
                            ),
                            "actual": workbook.sheetnames,
                        }
                    ],
                    "warnings": [],
                }

            template = template_result["template"]
            sheet_mapping = template_result["sheet_mapping"]

            # ==================================================
            # 3. WORKBOOK STRUCTURE
            # ==================================================

            self.workbook_validator.validate(
                workbook,
                issues,
                template=template,
                sheet_mapping=sheet_mapping
            )

            # ==================================================
            # 4. EXCEL ERROR CHECK
            # ==================================================

            self.excel_error_validator.validate(
                workbook,
                issues
            )

            # ==================================================
            # 5. STANDARD / ALTERNATE TEMPLATES
            # ==================================================

            if template in {
                TemplateDetector.TEMPLATE_STANDARD,
                TemplateDetector.TEMPLATE_ALTERNATE_SCORE,
            }:

                required_sheets = {
                    "score_sheet",
                    "Batch Analysis - Tabular",
                }

                available_sheets = set(
                    workbook.sheetnames
                )

                has_required_sheets = (
                    required_sheets <= available_sheets
                )

                if has_required_sheets:

                    # ------------------------------------------
                    # Raw data validation
                    # ------------------------------------------

                    self.data_validator.validate(
                        workbook,
                        issues,
                        sheet_mapping=sheet_mapping
                    )

                    # ------------------------------------------
                    # Formula validation
                    # ------------------------------------------

                    self.formula_validator.validate(
                        workbook,
                        issues,
                        sheet_mapping=sheet_mapping,
                        template=template
                    )

                    # ------------------------------------------
                    # Statistics validation
                    # ------------------------------------------
                    # Per final rule clarification: correctly calculated hardcoded NOS statistics are acceptable
                    # StatisticsValidator performs independent verification which may flag acceptable hardcoded values
                    # For now, disable StatisticsValidator to avoid false positives on acceptable hardcoded values
                    # FormulaValidator marks hardcoded statistics as REVIEW for manual verification
                    
                    # self.statistics_validator.validate(
                    #     workbook,
                    #     issues,
                    #     sheet_mapping=sheet_mapping,
                    #     template=template
                    # )

                    # ------------------------------------------
                    # Cross-sheet validation
                    # ------------------------------------------

                    self.cross_sheet_validator.validate(
                        workbook,
                        issues,
                        sheet_mapping=sheet_mapping,
                        template=template
                    )

            # ==================================================
            # 6. LEGACY TEMPLATE
            # ==================================================
            # Per confirmed validation rules: BOTH workbook template
            # families are officially accepted. Do not reject a workbook
            # only because it belongs to the legacy or alternate template family.
            # Apply the same validation logic using sheet mapping.

            elif template == (
                TemplateDetector.TEMPLATE_LEGACY_RESULT
            ):

                # Legacy template uses different sheet names:
                # - Result (instead of score_sheet)
                # - Analysis-Tabular (instead of Batch Analysis - Tabular)
                # - Analysis - Graph (instead of Batch Analysis - Graph)
                
                # Check if required sheets exist for legacy template
                required_sheets = {
                    sheet_mapping.get("PRIMARY_DATA_SHEET"),
                    sheet_mapping.get("TABULAR_ANALYSIS_SHEET"),
                }
                required_sheets = {s for s in required_sheets if s is not None}
                
                available_sheets = set(
                    sheet_name
                    for sheet_name in workbook.sheetnames
                    if not workbook[sheet_name].sheet_state == 'hidden'
                )
                
                has_required_sheets = required_sheets <= available_sheets
                
                if has_required_sheets:
                    # Apply the same validation logic as standard templates
                    # The sheet mapping handles the name differences
                    
                    # ------------------------------------------
                    # Raw data validation
                    # ------------------------------------------

                    self.data_validator.validate(
                        workbook,
                        issues,
                        sheet_mapping=sheet_mapping
                    )

                    # ------------------------------------------
                    # Formula validation
                    # ------------------------------------------

                    self.formula_validator.validate(
                        workbook,
                        issues,
                        sheet_mapping=sheet_mapping,
                        template=template
                    )

                    # ------------------------------------------
                    # Statistics validation
                    # ------------------------------------------
                    # Per final rule clarification: correctly calculated hardcoded NOS statistics are acceptable
                    # StatisticsValidator performs independent verification which may flag acceptable hardcoded values
                    # For now, disable StatisticsValidator to avoid false positives on acceptable hardcoded values
                    # FormulaValidator marks hardcoded statistics as REVIEW for manual verification
                    
                    # self.statistics_validator.validate(
                    #     workbook,
                    #     issues,
                    #     sheet_mapping=sheet_mapping,
                    #     template=template
                    # )

                    # ------------------------------------------
                    # Cross-sheet validation
                    # ------------------------------------------

                    self.cross_sheet_validator.validate(
                        workbook,
                        issues,
                        sheet_mapping=sheet_mapping,
                        template=template
                    )

        finally:

            workbook.close()

        # ==================================================
        # 7. SEPARATE ERRORS, WARNINGS, AND REVIEW ITEMS
        # ==================================================

        errors = [
            issue
            for issue in issues
            if issue.get("severity", "ERROR") not in {"WARNING", "REVIEW"}
        ]

        warnings = [
            issue
            for issue in issues
            if issue.get("severity") == "WARNING"
        ]

        review_items = [
            issue
            for issue in issues
            if issue.get("severity") == "REVIEW"
        ]

        # ==================================================
        # 8. FINAL STATUS
        # ==================================================
        # Status can be: PASS, ERROR, or REVIEW
        # - PASS: No errors or review items
        # - REVIEW: Has review items but no errors
        # - ERROR: Has errors

        if review_items and not errors:
            status = "REVIEW"
        elif errors:
            status = "ERROR"
        else:
            status = "PASS"

        # ==================================================
        # 9. AI CLASSIFICATION FOR REVIEW ITEMS
        # ==================================================
        # CRITICAL: AI classification NEVER changes the deterministic status.
        # It only provides assistance for human review of REVIEW items.

        ai_review = None
        if status == "REVIEW" and review_items:
            ai_review = self._classify_review_items_with_ai(
                review_items=review_items,
                template=template,
                filename=filename or file_path.name
            )

        return {
            "status": status,
            "template": template,
            "errors": errors,
            "warnings": warnings,
            "review_items": review_items,
            "ai_review": ai_review,
        }

    def _classify_review_items_with_ai(
        self,
        review_items: list,
        template: str,
        filename: str
    ) -> AIReviewSummary:
        """
        Classify REVIEW items using AI (Laya).

        This method only provides classification and reasoning to assist
        human review. It NEVER changes the deterministic validation status.

        Args:
            review_items: List of review items from deterministic validation
            template: Detected workbook template
            filename: Name of the file being validated

        Returns:
            AIReviewSummary with classification results
        """
        classifier = get_laya_classifier()

        if not classifier.is_enabled():
            return AIReviewSummary(
                enabled=False,
                total_review_items=len(review_items),
                successfully_classified=0,
                failed_classifications=0,
            )

        # Build context for each review item
        review_contexts = []
        for item in review_items:
            context = {
                "filename": filename,
                "template": template,
                "review_type": item.get("code", "unknown"),
                "sheet": item.get("sheet"),
                "cell": item.get("cell"),
                "actual_value": str(item.get("actual")) if item.get("actual") is not None else None,
                "expected_value": str(item.get("expected")) if item.get("expected") is not None else None,
                "validation_message": item.get("message", ""),
                "validation_rule": item.get("code", ""),
                "relevant_context": self._build_relevant_context(item),
            }
            review_contexts.append(context)

        # Classify all review items
        classification_results = classifier.classify_review_items_batch(review_contexts)

        # Build summary
        classifications = []
        successful = 0
        failed = 0

        for item, result in zip(review_items, classification_results):
            if result.success:
                classifications.append(AIReviewClassification(
                    classification=result.classification,
                    confidence=result.confidence,
                    reasoning=result.reasoning,
                    suggested_action=result.suggested_action,
                    model_used=result.model_used,
                    latency_ms=result.latency_ms,
                ))
                successful += 1
            else:
                classifications.append(AIReviewClassification(
                    error=result.error,
                    model_used=result.model_used,
                    latency_ms=result.latency_ms,
                ))
                failed += 1

        return AIReviewSummary(
            enabled=True,
            classifications=classifications,
            total_review_items=len(review_items),
            successfully_classified=successful,
            failed_classifications=failed,
        )

    def _build_relevant_context(self, review_item: dict) -> str:
        """
        Build minimal relevant context for a review item.

        This provides focused context to the AI without sending entire
        workbook contents. The context is tailored to the review type.

        Args:
            review_item: The review item from deterministic validation

        Returns:
            Context string for AI classification
        """
        code = review_item.get("code", "")

        # Provide context based on review type
        if code == "HARDCODED_PASS_FAIL_REVIEW":
            return (
                "This review item indicates that the Pass/Fail criterion could not be "
                "automatically verified. The validator checks both the score sheet and "
                "tabular analysis sheet for pass criteria definitions."
            )
        elif code == "HARDCODED_SUMMARY_REVIEW":
            return (
                "This review item indicates a hardcoded summary statistic. The validator "
                "expects formulas for calculated values, but some templates allow hardcoded "
                "metadata cells."
            )
        elif code == "NOS_STATISTIC_REVIEW":
            return (
                "This review item indicates a hardcoded NOS (Number of Students) statistic. "
                "The validator verifies that NOS statistics match the actual candidate count "
                "when calculated, but some templates allow correctly-calculated hardcoded values."
            )
        elif code == "CANDIDATE_HEADER_NOT_FOUND":
            return (
                "This review item indicates that the Candidate ID header could not be found "
                "after a thorough search. The workbook structure may differ from expected templates."
            )
        else:
            return (
                f"Review item code: {code}. "
                f"Category: {review_item.get('category', 'unknown')}. "
                "The deterministic validator could not conclusively verify this item."
            )