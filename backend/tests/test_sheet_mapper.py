"""
Tests for SheetMapper utility.
"""
import pytest
from openpyxl import Workbook
from pathlib import Path
import sys

# Add parent directory to path for imports
sys.path.insert(0, str(Path(__file__).parent.parent.parent))

from app.utils.sheet_mapper import SheetMapper


class TestSheetMapper:
    """Test SheetMapper functionality."""
    
    def test_standard_mapping(self):
        """Test standard template sheet mapping."""
        mapping = SheetMapper.get_mapping("STANDARD")
        
        assert mapping[SheetMapper.PRIMARY_DATA_SHEET] == "score_sheet"
        assert mapping[SheetMapper.TABULAR_ANALYSIS_SHEET] == "Batch Analysis - Tabular"
        assert mapping[SheetMapper.GRAPH_SHEET] == "Batch Analysis - Graph"
    
    def test_legacy_mapping(self):
        """Test legacy template sheet mapping."""
        mapping = SheetMapper.get_mapping("LEGACY_RESULT")
        
        assert mapping[SheetMapper.PRIMARY_DATA_SHEET] == "Result"
        assert mapping[SheetMapper.TABULAR_ANALYSIS_SHEET] == "Analysis-Tabular"
        assert mapping[SheetMapper.GRAPH_SHEET] == "Analysis - Graph"
    
    def test_alternate_mapping(self):
        """Test alternate template uses standard mapping."""
        mapping = SheetMapper.get_mapping("ALTERNATE_SCORE")
        
        assert mapping[SheetMapper.PRIMARY_DATA_SHEET] == "score_sheet"
        assert mapping[SheetMapper.TABULAR_ANALYSIS_SHEET] == "Batch Analysis - Tabular"
        assert mapping[SheetMapper.GRAPH_SHEET] == "Batch Analysis - Graph"
    
    def test_resolve_logical_name_standard(self):
        """Test resolving logical names with standard template."""
        wb = Workbook()
        
        # Create standard sheets
        score_ws = wb.active
        score_ws.title = "score_sheet"
        tabular_ws = wb.create_sheet("Batch Analysis - Tabular")
        graph_ws = wb.create_sheet("Batch Analysis - Graph")
        
        # Test resolution
        primary = SheetMapper.resolve_logical_name(wb, SheetMapper.PRIMARY_DATA_SHEET, "STANDARD")
        assert primary == "score_sheet"
        
        tabular = SheetMapper.resolve_logical_name(wb, SheetMapper.TABULAR_ANALYSIS_SHEET, "STANDARD")
        assert tabular == "Batch Analysis - Tabular"
        
        graph = SheetMapper.resolve_logical_name(wb, SheetMapper.GRAPH_SHEET, "STANDARD")
        assert graph == "Batch Analysis - Graph"
        
        wb.close()
    
    def test_resolve_logical_name_legacy(self):
        """Test resolving logical names with legacy template."""
        wb = Workbook()
        
        # Create legacy sheets
        result_ws = wb.active
        result_ws.title = "Result"
        analysis_ws = wb.create_sheet("Analysis-Tabular")
        graph_ws = wb.create_sheet("Analysis - Graph")
        
        # Test resolution
        primary = SheetMapper.resolve_logical_name(wb, SheetMapper.PRIMARY_DATA_SHEET, "LEGACY_RESULT")
        assert primary == "Result"
        
        tabular = SheetMapper.resolve_logical_name(wb, SheetMapper.TABULAR_ANALYSIS_SHEET, "LEGACY_RESULT")
        assert tabular == "Analysis-Tabular"
        
        graph = SheetMapper.resolve_logical_name(wb, SheetMapper.GRAPH_SHEET, "LEGACY_RESULT")
        assert graph == "Analysis - Graph"
        
        wb.close()
    
    def test_identify_primary_data_sheet_standard(self):
        """Test identifying primary data sheet for standard template."""
        wb = Workbook()
        
        # Create standard sheet with Candidate ID header
        score_ws = wb.active
        score_ws.title = "score_sheet"
        score_ws.cell(12, 1, "Candidate ID")
        score_ws.cell(12, 2, "Batch ID")
        
        # Create other sheet without Candidate ID
        other_ws = wb.create_sheet("Other Sheet")
        other_ws.cell(1, 1, "Some Data")
        
        # Test identification
        primary = SheetMapper.identify_primary_data_sheet(wb)
        assert primary == "score_sheet"
        
        wb.close()
    
    def test_identify_primary_data_sheet_legacy(self):
        """Test identifying primary data sheet for legacy template."""
        wb = Workbook()
        
        # Create legacy sheet with Candidate ID header
        result_ws = wb.active
        result_ws.title = "Result"
        result_ws.cell(10, 1, "Candidate ID")
        result_ws.cell(10, 2, "Batch ID")
        
        # Create other sheet without Candidate ID
        other_ws = wb.create_sheet("Analysis-Tabular")
        other_ws.cell(1, 1, "NOS SUMMARY")
        
        # Test identification
        primary = SheetMapper.identify_primary_data_sheet(wb)
        assert primary == "Result"
        
        wb.close()
    
    def test_identify_primary_data_sheet_ignores_hidden(self):
        """Test that hidden sheets are ignored in primary sheet identification."""
        wb = Workbook()
        
        # Create visible sheet with Candidate ID
        visible_ws = wb.active
        visible_ws.title = "visible_sheet"
        visible_ws.cell(12, 1, "Candidate ID")
        
        # Create hidden sheet with Candidate ID (should be ignored)
        hidden_ws = wb.create_sheet("hidden_sheet")
        hidden_ws.sheet_state = 'hidden'
        hidden_ws.cell(12, 1, "Candidate ID")
        
        # Test identification
        primary = SheetMapper.identify_primary_data_sheet(wb)
        assert primary == "visible_sheet"
        
        wb.close()
    
    def test_identify_tabular_analysis_sheet(self):
        """Test identifying tabular analysis sheet."""
        wb = Workbook()
        
        # Create tabular sheet with NOS SUMMARY
        tabular_ws = wb.active
        tabular_ws.title = "Batch Analysis - Tabular"
        tabular_ws.cell(10, 1, "NOS SUMMARY")
        tabular_ws.cell(12, 1, "SSC/N8417")
        
        # Create other sheet without NOS
        other_ws = wb.create_sheet("Other Sheet")
        other_ws.cell(1, 1, "Some Data")
        
        # Test identification
        tabular = SheetMapper.identify_tabular_analysis_sheet(wb)
        assert tabular == "Batch Analysis - Tabular"
        
        wb.close()
    
    def test_identify_tabular_analysis_sheet_ignores_hidden(self):
        """Test that hidden sheets are ignored in tabular sheet identification."""
        wb = Workbook()
        
        # Create visible sheet with NOS
        visible_ws = wb.active
        visible_ws.title = "visible_sheet"
        visible_ws.cell(10, 1, "NOS SUMMARY")
        visible_ws.cell(12, 1, "SSC/N8417")
        
        # Create hidden sheet with NOS (should be ignored)
        hidden_ws = wb.create_sheet("hidden_sheet")
        hidden_ws.sheet_state = 'hidden'
        hidden_ws.cell(10, 1, "NOS SUMMARY")
        hidden_ws.cell(12, 1, "SSC/N8417")
        
        # Test identification
        tabular = SheetMapper.identify_tabular_analysis_sheet(wb)
        assert tabular == "visible_sheet"
        
        wb.close()