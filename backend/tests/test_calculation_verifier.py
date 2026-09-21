"""
Tests for CalculationVerifier utility.
"""
import pytest
from openpyxl import Workbook
from pathlib import Path
import sys

# Add parent directory to path for imports
sys.path.insert(0, str(Path(__file__).parent.parent.parent))

from backend.app.utils.calculation_verifier import CalculationVerifier


class TestCalculationVerifier:
    """Test CalculationVerifier functionality."""
    
    def test_verify_enrolled_count(self):
        """Test independent calculation of enrolled count."""
        wb = Workbook()
        ws = wb.active
        
        # Header at row 10
        ws.cell(10, 1, "Candidate ID")
        ws.cell(10, 2, "Batch ID")
        
        # 5 candidate rows
        for i in range(11, 16):
            ws.cell(i, 1, f"CAND-{i}")
            ws.cell(i, 2, "BATCH-001")
        
        columns = {"candidate_id": 1, "batch_id": 2}
        result = CalculationVerifier.verify_enrolled_count(ws, 10, columns)
        
        assert result["verification_possible"] == True
        assert result["expected"] == 5
        assert result["status"] == "CALCULATED"
        
        wb.close()
    
    def test_verify_appeared_count(self):
        """Test independent calculation of appeared count."""
        wb = Workbook()
        ws = wb.active
        
        # Header at row 10
        ws.cell(10, 1, "Candidate ID")
        ws.cell(10, 2, "Batch ID")
        ws.cell(10, 3, "SSC/N8417 - Score")
        
        # 5 candidate rows, but only 3 have scores
        ws.cell(11, 1, "CAND-11")
        ws.cell(11, 2, "BATCH-001")
        ws.cell(11, 3, 85)
        
        ws.cell(12, 1, "CAND-12")
        ws.cell(12, 2, "BATCH-001")
        ws.cell(12, 3, 92)
        
        ws.cell(13, 1, "CAND-13")
        ws.cell(13, 2, "BATCH-001")
        ws.cell(13, 3, None)  # No score
        
        ws.cell(14, 1, "CAND-14")
        ws.cell(14, 2, "BATCH-001")
        ws.cell(14, 3, 78)
        
        ws.cell(15, 1, "CAND-15")
        ws.cell(15, 2, "BATCH-001")
        ws.cell(15, 3, None)  # No score
        
        columns = {"candidate_id": 1, "batch_id": 2}
        score_columns = [{"score_column": 3, "nos": "SSC/N8417", "max_score": 100}]
        
        result = CalculationVerifier.verify_appeared_count(ws, 10, columns, score_columns)
        
        assert result["verification_possible"] == True
        assert result["expected"] == 3  # Only 3 have scores
        assert result["status"] == "CALCULATED"
        
        wb.close()
    
    def test_verify_passed_count(self):
        """Test independent calculation of passed count."""
        wb = Workbook()
        ws = wb.active
        
        # Header at row 10
        ws.cell(10, 1, "Candidate ID")
        ws.cell(10, 2, "Batch ID")
        ws.cell(10, 3, "SSC/N8417 - Score")
        ws.cell(10, 4, "Pass/Fail")
        
        # 4 candidate rows with different results
        ws.cell(11, 1, "CAND-11")
        ws.cell(11, 2, "BATCH-001")
        ws.cell(11, 3, 85)  # Pass (>= 80)
        ws.cell(11, 4, "=IF(C11/100*100>=80,\"Pass\",\"Fail\")")
        
        ws.cell(12, 1, "CAND-12")
        ws.cell(12, 2, "BATCH-001")
        ws.cell(12, 3, 92)  # Pass
        ws.cell(12, 4, "=IF(C12/100*100>=80,\"Pass\",\"Fail\")")
        
        ws.cell(13, 1, "CAND-13")
        ws.cell(13, 2, "BATCH-001")
        ws.cell(13, 3, 75)  # Fail
        ws.cell(13, 4, "=IF(C13/100*100>=80,\"Pass\",\"Fail\")")
        
        ws.cell(14, 1, "CAND-14")
        ws.cell(14, 2, "BATCH-001")
        ws.cell(14, 3, None)  # No score - not appeared
        ws.cell(14, 4, None)
        
        columns = {"candidate_id": 1, "batch_id": 2, "pass_fail": 4}
        score_columns = [{"score_column": 3, "nos": "SSC/N8417", "max_score": 100}]
        candidate_rows = [11, 12, 13, 14]
        
        result = CalculationVerifier.verify_passed_count(ws, 10, columns, score_columns, 0.8)
        
        assert result["verification_possible"] == True
        assert result["expected"] == 2  # Only 2 passed (candidates 11 and 12)
        assert result["status"] == "CALCULATED"
        
        wb.close()
    
    def test_verify_gender_counts(self):
        """Test independent calculation of gender counts."""
        wb = Workbook()
        ws = wb.active
        
        # Header at row 10
        ws.cell(10, 1, "Candidate ID")
        ws.cell(10, 2, "Gender")
        
        # 5 candidate rows with different genders
        ws.cell(11, 1, "CAND-11")
        ws.cell(11, 2, "Male")
        
        ws.cell(12, 1, "CAND-12")
        ws.cell(12, 2, "Female")
        
        ws.cell(13, 1, "CAND-13")
        ws.cell(13, 2, "M")
        
        ws.cell(14, 1, "CAND-14")
        ws.cell(14, 2, "F")
        
        ws.cell(15, 1, "CAND-15")
        ws.cell(15, 2, "NA")  # Not counted
        
        columns = {"candidate_id": 1, "gender": 2}
        
        result = CalculationVerifier.verify_gender_counts(ws, 10, columns)
        
        assert result["verification_possible"] == True
        assert result["male_count"] == 2  # Male and M
        assert result["female_count"] == 2  # Female and F
        assert result["status"] == "CALCULATED"
        
        wb.close()
    
    def test_verify_nos_statistics(self):
        """Test independent calculation of NOS statistics."""
        wb = Workbook()
        ws = wb.active
        
        # Header at row 10
        ws.cell(10, 1, "Candidate ID")
        ws.cell(10, 2, "SSC/N8417 - Score")
        
        # 5 candidate rows with scores
        scores = [85, 92, 78, 88, 95]
        for i, score in enumerate(scores, start=11):
            ws.cell(i, 1, f"CAND-{i}")
            ws.cell(i, 2, score)
        
        columns = {"candidate_id": 1}
        score_columns = [{"score_column": 2, "nos": "SSC/N8417", "max_score": 100}]
        candidate_rows = [11, 12, 13, 14, 15]
        
        result = CalculationVerifier.verify_nos_statistics(ws, 10, score_columns, candidate_rows, "SSC/N8417", 100)
        
        assert result["verification_possible"] == True
        assert result["status"] == "CALCULATED"
        assert result["mean"] == sum(scores) / len(scores)
        assert result["min"] == min(scores)
        assert result["max"] == max(scores)
        assert result["pass_percentage"] == 0.8  # 4 out of 5 passed (78 failed)
        
        wb.close()
    
    def test_compare_values_match(self):
        """Test value comparison with matching values."""
        result = CalculationVerifier.compare_values(100, 100, 0, "test_field")
        
        assert result["match"] == True
        assert result["status"] == "PASS"
        assert result["difference"] == 0
    
    def test_compare_values_mismatch(self):
        """Test value comparison with mismatching values."""
        result = CalculationVerifier.compare_values(100, 95, 0, "test_field")
        
        assert result["match"] == False
        assert result["status"] == "FAIL"
        assert result["difference"] == 5
    
    def test_compare_values_within_tolerance(self):
        """Test value comparison within tolerance."""
        result = CalculationVerifier.compare_values(100, 100.005, 0.01, "test_field")
        
        assert result["match"] == True
        assert result["status"] == "PASS"
    
    def test_compare_values_review_expected_none(self):
        """Test value comparison when expected cannot be calculated."""
        result = CalculationVerifier.compare_values(None, 100, 0, "test_field")
        
        assert result["status"] == "REVIEW"
        assert "reason" in result
    
    def test_compare_values_review_actual_none(self):
        """Test value comparison when actual is missing."""
        result = CalculationVerifier.compare_values(100, None, 0, "test_field")
        
        assert result["status"] == "REVIEW"
        assert "reason" in result
    
    def test_normalize_to_numeric(self):
        """Test normalization of various value formats."""
        # Test percentage
        assert CalculationVerifier.normalize_to_numeric("80%") == 0.8
        assert CalculationVerifier.normalize_to_numeric("75.5%") == 0.755
        
        # Test decimal
        assert CalculationVerifier.normalize_to_numeric(0.8) == 0.8
        assert CalculationVerifier.normalize_to_numeric(85) == 85.0
        
        # Test text percentage
        assert CalculationVerifier.normalize_to_numeric("80 percent") == 0.8
        
        # Test invalid
        assert CalculationVerifier.normalize_to_numeric("invalid") is None
        assert CalculationVerifier.normalize_to_numeric(None) is None