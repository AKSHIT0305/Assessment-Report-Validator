"""
Tests for WorkbookAnalyzer utility.
"""
import pytest
from openpyxl import Workbook
from pathlib import Path
import sys

# Add parent directory to path for imports
sys.path.insert(0, str(Path(__file__).parent.parent.parent))

from app.utils.workbook_analyzer import WorkbookAnalyzer


class TestWorkbookAnalyzer:
    """Test WorkbookAnalyzer functionality."""
    
    def test_find_header_row(self):
        """Test finding header row with Candidate ID."""
        wb = Workbook()
        ws = wb.active
        
        # Add Candidate ID header at row 15
        ws.cell(15, 3, "Candidate ID")
        ws.cell(15, 4, "Batch ID")
        
        # Add data at row 16
        ws.cell(16, 3, "CAND-001")
        ws.cell(16, 4, "BATCH-001")
        
        header_row = WorkbookAnalyzer.find_header_row(ws)
        assert header_row == 15
        
        wb.close()
    
    def test_find_header_row_variations(self):
        """Test finding header row with different variations."""
        test_variations = [
            "Candidate ID",
            "CandidateID",
            "candidate id",
            "CANDIDATE ID",
        ]
        
        for i, variation in enumerate(test_variations):
            wb = Workbook()
            ws = wb.active
            
            # Add variation at different row
            row = 10 + i
            ws.cell(row, 1, variation)
            ws.cell(row, 2, "Batch ID")
            ws.cell(row + 1, 1, "CAND-001")
            
            header_row = WorkbookAnalyzer.find_header_row(ws)
            assert header_row == row, f"Failed for variation: {variation}"
            
            wb.close()
    
    def test_find_data_region(self):
        """Test finding data region."""
        wb = Workbook()
        ws = wb.active
        
        # Header at row 10
        ws.cell(10, 1, "Candidate ID")
        ws.cell(10, 2, "Batch ID")
        
        # Data rows 11-15
        for i in range(11, 16):
            ws.cell(i, 1, f"CAND-{i}")
            ws.cell(i, 2, f"BATCH-{i}")
        
        # Empty row at 16
        ws.cell(17, 1, "CAND-017")
        ws.cell(17, 2, "BATCH-017")
        
        data_region = WorkbookAnalyzer.find_data_region(ws, 10)
        assert data_region == (11, 17)
        
        wb.close()
    
    def test_detect_columns(self):
        """Test column detection."""
        wb = Workbook()
        ws = wb.active
        
        # Header row with columns in different order
        ws.cell(12, 1, "Gender")
        ws.cell(12, 2, "Candidate ID")
        ws.cell(12, 3, "Batch ID")
        ws.cell(12, 4, "Assessment Date")
        
        columns = WorkbookAnalyzer.detect_columns(ws, 12)
        
        assert columns["candidate_id"] == 2
        assert columns["batch_id"] == 3
        assert columns["gender"] == 1
        assert columns["assessment_date"] == 4
        
        wb.close()
    
    def test_find_summary_section(self):
        """Test finding summary section dynamically."""
        wb = Workbook()
        ws = wb.active
        
        # Add summary section at different location
        ws.cell(20, 1, "BATCH SUMMARY")
        ws.cell(21, 1, "Enrolled")
        ws.cell(21, 2, 100)
        ws.cell(22, 1, "Appeared")
        ws.cell(22, 2, 95)
        ws.cell(23, 1, "Passed")
        ws.cell(23, 2, 80)
        
        summary_cells = WorkbookAnalyzer.find_summary_section(ws)
        
        assert "enrolled" in summary_cells
        assert "appeared" in summary_cells
        assert "passed" in summary_cells
        
        wb.close()
    
    def test_find_batch_id_cell(self):
        """Test finding batch ID cell dynamically."""
        wb = Workbook()
        ws = wb.active
        
        # Add batch ID at different location
        ws.cell(5, 3, "Batch ID")
        ws.cell(5, 4, "TEST-12345")
        
        batch_id_cell = WorkbookAnalyzer.find_batch_id_cell(ws)
        assert batch_id_cell == "D5"
        
        wb.close()
    
    def test_find_nos_section(self):
        """Test finding NOS section dynamically."""
        wb = Workbook()
        ws = wb.active
        
        # Add NOS codes at different location
        ws.cell(25, 1, "SSC/N8417")
        ws.cell(26, 1, "MEP/N2601")
        ws.cell(27, 1, "WEAK PC - Performance")
        
        nos_section = WorkbookAnalyzer.find_nos_section(ws)
        
        assert nos_section is not None
        assert nos_section["start_row"] == 25
        assert nos_section["end_row"] == 27
        assert len(nos_section["nos_codes"]) == 3
        
        wb.close()
    
    def test_detect_pass_criteria(self):
        """Test detecting pass criteria."""
        wb = Workbook()
        ws = wb.active
        
        # Add pass criteria in different formats
        ws.cell(5, 1, "Pass Criteria")
        ws.cell(5, 2, "80%")
        
        criteria = WorkbookAnalyzer.detect_pass_criteria(ws)
        assert criteria == 0.8
        
        wb.close()
    
    def test_parse_pass_criteria(self):
        """Test parsing pass criteria from various formats."""
        # Test percentage format
        assert WorkbookAnalyzer.parse_pass_criteria("80%") == 0.8
        assert WorkbookAnalyzer.parse_pass_criteria("75.5%") == 0.755
        
        # Test fraction format
        assert WorkbookAnalyzer.parse_pass_criteria("80/100") == 0.8
        assert WorkbookAnalyzer.parse_pass_criteria("40/50") == 0.8
        
        # Test decimal format
        assert WorkbookAnalyzer.parse_pass_criteria(0.8) == 0.8
        assert WorkbookAnalyzer.parse_pass_criteria(0.75) == 0.75
        
        # Test text format
        assert WorkbookAnalyzer.parse_pass_criteria("80 percent") == 0.8
        
        # Test invalid format
        assert WorkbookAnalyzer.parse_pass_criteria("invalid") is None
    
    def test_normalize_text(self):
        """Test text normalization."""
        assert WorkbookAnalyzer.normalize_text("Candidate ID") == "candidate id"
        assert WorkbookAnalyzer.normalize_text("CandidateID") == "candidateid"
        assert WorkbookAnalyzer.normalize_text("Batch_ID") == "batch id"
        assert WorkbookAnalyzer.normalize_text("  Extra  Spaces  ") == "extra spaces"